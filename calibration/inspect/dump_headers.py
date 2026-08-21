from pathlib import Path
import openpyxl
REPO = Path(__file__).resolve().parents[2]
wb = openpyxl.load_workbook(REPO/"data/raw/T1_MVrawDataR2_2025_12_completed.xlsx",
                            read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
for i, h in enumerate(hdr):
    if h is None: continue
    s = str(h).replace("\n", " ").strip()
    if any(k in s.lower() for k in ("header", "hcw", "flow", "fls", "rt", "kw/", "load", "time", "date")):
        print(f"  col {i:3d}  {s}")
print("\n--- audit sheet ---")
a = wb["Calculation_Audit"]
for r in a.iter_rows(values_only=True):
    vals = [str(v)[:70] for v in r if v is not None]
    if vals: print("  " + " | ".join(vals))
wb.close()
